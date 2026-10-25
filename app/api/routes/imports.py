"""Safe, tenant-scoped spreadsheet migrations.

The API deliberately receives the file twice (preview and commit).  A preview is
never trusted as an authorization token: the commit reparses and revalidates the
same file inside one database transaction.
"""

from __future__ import annotations

import csv
from hashlib import sha256
import io
import re
import unicodedata
from collections import defaultdict
from datetime import date, datetime, time, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from typing import Any, Callable

from fastapi import APIRouter, Depends, File, HTTPException, Request, UploadFile, status
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.db.session import get_db
from app.models.entities import (
    AccountPayable,
    AccountReceivable,
    Category,
    Customer,
    Product,
    Purchase,
    PurchaseItem,
    Sale,
    SaleItem,
    Supplier,
    UnitOfMeasure,
    Warehouse,
)
from app.security.auth import Actor, require_owner_or_admin, require_second_factor
from app.services.accounting import post_entry, split_tax
from app.services.audit import record
from app.services.idempotency import claim as claim_idempotency, complete as complete_idempotency, replay_response
from app.services.operational_context import resolve_warehouse_context
from app.services.stock import change_stock
from app.services.taxes import resolve_default_tax_rate
from app.services.cost_centers import resolve_cost_center

router = APIRouter(prefix="/imports", tags=["imports"])

MAX_FILE_SIZE = 5 * 1024 * 1024
MAX_ROWS = 2_000

TEMPLATES: dict[str, dict[str, Any]] = {
    "products": {
        "label": "Productos",
        "columns": ["sku", "name", "category", "unit", "sale_price", "average_cost"],
        "example": ["SKU-001", "Producto de ejemplo", "General", "unit", "100.00", "80.00"],
    },
    "customers": {
        "label": "Clientes",
        "columns": ["code", "name", "tax_id", "email", "phone", "credit_limit"],
        "example": ["CLI-001", "Cliente de ejemplo", "12345678", "contacto@cliente.com", "999000000", "0"],
    },
    "suppliers": {
        "label": "Proveedores",
        "columns": ["code", "name", "tax_id", "email", "phone", "address"],
        "example": ["PRV-001", "Proveedor de ejemplo", "20123456789", "ventas@proveedor.com", "999000000", "Dirección fiscal"],
    },
    "warehouses": {
        "label": "Almacenes",
        "columns": ["code", "name", "address"],
        "example": ["ALM-001", "Almacén principal", "Dirección del almacén"],
    },
    "inventory": {
        "label": "Ajustes de inventario",
        "columns": ["number", "sku", "warehouse_code", "quantity", "reason", "reference", "note"],
        "example": ["AJ-20261019-001", "SKU-001", "ALM-001", "25", "Carga inicial", "MIG-001", "Existencia inicial"],
    },
    "sales": {
        "label": "Ventas",
        "columns": ["number", "customer_tax_id", "customer_name", "warehouse_code", "product_sku", "quantity", "unit_price", "due_date", "note"],
        "example": ["V-20261019-001", "12345678", "Cliente de ejemplo", "ALM-001", "SKU-001", "2", "100.00", "2026-10-19", "Migración histórica"],
    },
    "purchases": {
        "label": "Compras",
        "columns": ["number", "supplier_tax_id", "supplier_name", "warehouse_code", "product_sku", "quantity", "unit_cost", "due_date", "note"],
        "example": ["C-20261019-001", "20123456789", "Proveedor de ejemplo", "ALM-001", "SKU-001", "10", "80.00", "2026-10-19", "Migración histórica"],
    },
}


def _header(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]+", "_", text.lower()).strip("_")


def _text(value: Any) -> str:
    return str(value or "").strip()


def _decimal(value: Any, field: str, *, whole: bool = False, positive: bool = False) -> Decimal:
    raw = _text(value).replace(" ", "").replace(",", ".")
    try:
        number = Decimal(raw)
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"{field} debe ser numérico") from exc
    if whole and number != number.to_integral_value():
        raise ValueError(f"{field} debe ser un número entero")
    if positive and number <= 0:
        raise ValueError(f"{field} debe ser mayor que cero")
    return number.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP) if not whole else number


def _date(value: Any, field: str) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(_text(value))
    except ValueError as exc:
        raise ValueError(f"{field} debe usar el formato AAAA-MM-DD") from exc


async def _rows_from_upload(upload: UploadFile) -> list[dict[str, Any]]:
    filename = (upload.filename or "").lower()
    if not filename.endswith((".csv", ".xlsx")):
        raise HTTPException(status_code=422, detail="Carga una plantilla CSV o XLSX")
    content = await upload.read(MAX_FILE_SIZE + 1)
    if not content or len(content) > MAX_FILE_SIZE:
        raise HTTPException(status_code=422, detail="La plantilla debe contener datos y pesar como máximo 5 MB")
    if filename.endswith(".csv"):
        try:
            reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
            rows = [{_header(key): value for key, value in row.items()} for row in reader]
        except UnicodeDecodeError as exc:
            raise HTTPException(status_code=422, detail="El CSV debe estar codificado en UTF-8") from exc
    else:
        try:
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            sheet = book.active
            values = sheet.iter_rows(values_only=True)
            headers = next(values, None)
            if not headers:
                raise ValueError("Sin encabezados")
            normalized_headers = [_header(header) for header in headers]
            rows = [
                {normalized_headers[index]: value for index, value in enumerate(row) if normalized_headers[index]}
                for row in values
                if any(value not in (None, "") for value in row)
            ]
        except Exception as exc:
            raise HTTPException(status_code=422, detail="No se pudo leer el archivo XLSX") from exc
    if not rows:
        raise HTTPException(status_code=422, detail="La plantilla no contiene filas de datos")
    if len(rows) > MAX_ROWS:
        raise HTTPException(status_code=422, detail=f"La plantilla admite hasta {MAX_ROWS} filas por importación")
    return rows


def _required(module: str, rows: list[dict[str, Any]]) -> list[str]:
    columns = TEMPLATES[module]["columns"]
    present = set(rows[0])
    return [column for column in columns if column not in present]


def _validate_master(module: str, row: dict[str, Any]) -> dict[str, Any]:
    if module == "products":
        sku, name = _text(row.get("sku")), _text(row.get("name"))
        if not sku or not name:
            raise ValueError("sku y name son obligatorios")
        return {"sku": sku, "name": name, "category": _text(row.get("category")), "unit": _text(row.get("unit")) or "unit", "sale_price": _decimal(row.get("sale_price") or 0, "sale_price"), "average_cost": _decimal(row.get("average_cost") or 0, "average_cost")}
    if module == "customers":
        code, name = _text(row.get("code")), _text(row.get("name"))
        tax_id = _text(row.get("tax_id"))
        if not code or not name:
            raise ValueError("code y name son obligatorios")
        if tax_id and (not tax_id.isdigit() or len(tax_id) not in {8, 11}):
            raise ValueError("tax_id debe ser DNI de 8 o RUC de 11 dígitos")
        return {"code": code, "name": name, "tax_id": tax_id or None, "email": _text(row.get("email")) or None, "phone": _text(row.get("phone")) or None, "credit_limit": _decimal(row.get("credit_limit") or 0, "credit_limit")}
    if module == "suppliers":
        code, name, tax_id = _text(row.get("code")), _text(row.get("name")), _text(row.get("tax_id"))
        if not code or not name:
            raise ValueError("code y name son obligatorios")
        if tax_id and (not tax_id.isdigit() or len(tax_id) != 11):
            raise ValueError("tax_id debe ser un RUC de 11 dígitos")
        return {"code": code, "name": name, "tax_id": tax_id or None, "email": _text(row.get("email")) or None, "phone": _text(row.get("phone")) or None, "address": _text(row.get("address")) or None}
    if module == "warehouses":
        code, name = _text(row.get("code")), _text(row.get("name"))
        if not code or not name:
            raise ValueError("code y name son obligatorios")
        return {"code": code, "name": name, "address": _text(row.get("address")) or None}
    raise ValueError("Módulo de migración no reconocido")


def _validate_transaction(module: str, row: dict[str, Any]) -> dict[str, Any]:
    quantity = _decimal(row.get("quantity"), "quantity", whole=True, positive=module != "inventory")
    if module == "inventory" and quantity == 0:
        raise ValueError("quantity no puede ser cero")
    # Inventory adjustments use the concise `sku` header, while sales and
    # purchases retain `product_sku` for clarity in multi-document sheets.
    product_sku = _text(row.get("sku" if module == "inventory" else "product_sku"))
    common = {"number": _text(row.get("number")), "warehouse_code": _text(row.get("warehouse_code")), "product_sku": product_sku, "quantity": quantity, "due_date": _date(row.get("due_date"), "due_date"), "note": _text(row.get("note")) or None}
    if not common["number"] or not common["warehouse_code"] or not common["product_sku"]:
        raise ValueError("number, warehouse_code y product_sku son obligatorios")
    if module == "inventory":
        reason = _text(row.get("reason"))
        if not reason:
            raise ValueError("reason es obligatorio")
        return {**common, "reason": reason, "reference": _text(row.get("reference")) or common["number"]}
    if module == "sales":
        tax_id, name = _text(row.get("customer_tax_id")), _text(row.get("customer_name"))
        if not tax_id.isdigit() or len(tax_id) not in {8, 11} or not name:
            raise ValueError("customer_tax_id (DNI/RUC) y customer_name son obligatorios")
        return {**common, "counterpart_tax_id": tax_id, "counterpart_name": name, "unit_value": _decimal(row.get("unit_price"), "unit_price")}
    if module == "purchases":
        tax_id, name = _text(row.get("supplier_tax_id")), _text(row.get("supplier_name"))
        if not tax_id.isdigit() or len(tax_id) != 11 or not name:
            raise ValueError("supplier_tax_id (RUC) y supplier_name son obligatorios")
        return {**common, "counterpart_tax_id": tax_id, "counterpart_name": name, "unit_value": _decimal(row.get("unit_cost"), "unit_cost")}
    raise ValueError("Módulo de migración no reconocido")


def _validate_rows(module: str, rows: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    missing = _required(module, rows)
    if missing:
        return [], [{"row": 1, "message": f"Faltan columnas: {', '.join(missing)}"}]
    valid, errors = [], []
    validator: Callable[[str, dict[str, Any]], dict[str, Any]] = _validate_master if module in {"products", "customers", "suppliers", "warehouses"} else _validate_transaction
    for index, row in enumerate(rows, start=2):
        try:
            valid.append(validator(module, row))
        except ValueError as exc:
            errors.append({"row": index, "message": str(exc)})
    if module in {"products", "customers", "suppliers", "warehouses"}:
        key = "sku" if module == "products" else "code"
        seen: set[str] = set()
        for index, row in enumerate(valid, start=2):
            value = row[key].casefold()
            if value in seen:
                errors.append({"row": index, "message": f"{key} está repetido dentro de la plantilla"})
            seen.add(value)
    return valid, errors


def _check_database_conflicts(module: str, rows: list[dict[str, Any]], actor: Actor, db: Session) -> list[str]:
    if module == "products":
        existing = set(db.scalars(select(Product.sku).where(Product.tenant_id == actor.tenant_id, Product.sku.in_([row["sku"] for row in rows]))).all())
        return [f"SKU ya existe: {value}" for value in existing]
    if module == "customers":
        existing = set(db.scalars(select(Customer.code).where(Customer.tenant_id == actor.tenant_id, Customer.code.in_([row["code"] for row in rows]))).all())
        return [f"Código de cliente ya existe: {value}" for value in existing]
    if module == "suppliers":
        existing = set(db.scalars(select(Supplier.code).where(Supplier.tenant_id == actor.tenant_id, Supplier.code.in_([row["code"] for row in rows]))).all())
        return [f"Código de proveedor ya existe: {value}" for value in existing]
    if module == "warehouses":
        existing = set(db.scalars(select(Warehouse.code).where(Warehouse.tenant_id == actor.tenant_id, Warehouse.code.in_([row["code"] for row in rows]))).all())
        return [f"Código de almacén ya existe: {value}" for value in existing]
    if module == "sales":
        numbers = {row["number"] for row in rows}
        existing = set(db.scalars(select(Sale.number).where(Sale.tenant_id == actor.tenant_id, Sale.number.in_(numbers))).all())
        return [f"Venta ya existe: {value}" for value in existing]
    if module == "purchases":
        numbers = {row["number"] for row in rows}
        existing = set(db.scalars(select(Purchase.number).where(Purchase.tenant_id == actor.tenant_id, Purchase.number.in_(numbers))).all())
        return [f"Compra ya existe: {value}" for value in existing]
    return []


def _reference_maps(actor: Actor, db: Session) -> tuple[dict[str, Product], dict[str, Warehouse]]:
    products = {item.sku: item for item in db.scalars(select(Product).where(Product.tenant_id == actor.tenant_id, Product.is_active.is_(True))).all()}
    warehouses = {item.code: item for item in db.scalars(select(Warehouse).where(Warehouse.tenant_id == actor.tenant_id, Warehouse.is_active.is_(True))).all()}
    return products, warehouses


def _transaction_references(module: str, rows: list[dict[str, Any]], actor: Actor, db: Session) -> list[str]:
    products, warehouses = _reference_maps(actor, db)
    errors: list[str] = []
    for row in rows:
        if row["product_sku"] not in products:
            errors.append(f"Producto no encontrado o inactivo: {row['product_sku']}")
        if row["warehouse_code"] not in warehouses:
            errors.append(f"Almacén no encontrado o inactivo: {row['warehouse_code']}")
    return list(dict.fromkeys(errors))


def _persist_master(module: str, rows: list[dict[str, Any]], actor: Actor, db: Session) -> int:
    for row in rows:
        if module == "products":
            category = None
            if row["category"]:
                category = db.scalar(select(Category).where(Category.tenant_id == actor.tenant_id, Category.name == row["category"]))
                if category is None:
                    category = Category(tenant_id=actor.tenant_id, name=row["category"])
                    db.add(category); db.flush()
            unit_value = row["unit"].strip()
            unit = db.scalar(select(UnitOfMeasure).where(
                UnitOfMeasure.tenant_id == actor.tenant_id,
                UnitOfMeasure.is_active.is_(True),
                (UnitOfMeasure.code.ilike(unit_value)) | (UnitOfMeasure.name.ilike(unit_value)),
            ))
            if unit is None:
                raise HTTPException(status_code=422, detail=f"Unidad no encontrada o inactiva: {unit_value}. Créala en Configuración antes de importar.")
            entity = Product(tenant_id=actor.tenant_id, sku=row["sku"], name=row["name"], category_id=category.id if category else None, unit_id=unit.id, unit=unit.code, unit_code=unit.code, unit_name=unit.name, sale_price=row["sale_price"], average_cost=row["average_cost"])
        elif module == "customers":
            entity = Customer(tenant_id=actor.tenant_id, **row)
        elif module == "suppliers":
            entity = Supplier(tenant_id=actor.tenant_id, **row)
        else:
            entity = Warehouse(tenant_id=actor.tenant_id, **row)
        db.add(entity); db.flush()
        record(db, tenant_id=actor.tenant_id, actor_user_id=actor.user_id, action="import", module=module, entity=module.rstrip("s"), entity_id=entity.id)
    return len(rows)


def _get_or_create_counterpart(module: str, row: dict[str, Any], actor: Actor, db: Session) -> Customer | Supplier:
    if module == "sales":
        entity = db.scalar(select(Customer).where(Customer.tenant_id == actor.tenant_id, Customer.tax_id == row["counterpart_tax_id"]))
        if entity is None:
            entity = Customer(tenant_id=actor.tenant_id, code=f"CLI-{row['counterpart_tax_id']}", name=row["counterpart_name"], tax_id=row["counterpart_tax_id"])
            db.add(entity); db.flush()
        return entity
    entity = db.scalar(select(Supplier).where(Supplier.tenant_id == actor.tenant_id, Supplier.tax_id == row["counterpart_tax_id"]))
    if entity is None:
        entity = Supplier(tenant_id=actor.tenant_id, code=f"PRV-{row['counterpart_tax_id']}", name=row["counterpart_name"], tax_id=row["counterpart_tax_id"])
        db.add(entity); db.flush()
    return entity


def _persist_transactions(module: str, rows: list[dict[str, Any]], actor: Actor, db: Session) -> int:
    products, warehouses = _reference_maps(actor, db)
    if module == "inventory":
        for row in rows:
            product, warehouse = products[row["product_sku"]], warehouses[row["warehouse_code"]]
            change_stock(db, tenant_id=actor.tenant_id, product_id=product.id, warehouse_id=warehouse.id, quantity=row["quantity"], movement_type="adjustment", reference=row["reference"], note=row["note"] or row["reason"])
            record(db, tenant_id=actor.tenant_id, actor_user_id=actor.user_id, action="import_adjustment", module="inventory", entity="inventory_movement", detail=row["reason"])
        return len(rows)
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        grouped[row["number"]].append(row)
    for number, lines in grouped.items():
        first, warehouse = lines[0], warehouses[lines[0]["warehouse_code"]]
        if any(line["warehouse_code"] != first["warehouse_code"] for line in lines):
            raise HTTPException(status_code=422, detail=f"El documento {number} mezcla almacenes")
        if len({line["product_sku"] for line in lines}) != len(lines):
            raise HTTPException(status_code=422, detail=f"El documento {number} repite un SKU")
        counterpart = _get_or_create_counterpart(module, first, actor, db)
        total = sum((line["quantity"] * line["unit_value"] for line in lines), Decimal("0"))
        tax_rate = resolve_default_tax_rate(
            db, tenant_id=actor.tenant_id, document_date=date.today(),
        )
        taxable_amount, tax_amount = split_tax(total, tax_rate.rate)
        cost_center = resolve_cost_center(
            db, tenant_id=actor.tenant_id, cost_center_id=None,
        )
        operational_context = resolve_warehouse_context(
            db, tenant_id=actor.tenant_id, warehouse_id=warehouse.id,
        )
        if module == "sales":
            sale = Sale(
                tenant_id=actor.tenant_id, number=number, customer_id=counterpart.id,
                warehouse_id=warehouse.id,
                legal_entity_id=operational_context.legal_entity.id if operational_context.legal_entity else None,
                branch_id=operational_context.branch.id if operational_context.branch else None,
                status="posted", total=total, note=first["note"],
                cost_center_id=cost_center.id, cost_center_code=cost_center.code,
                cost_center_name=cost_center.name,
                tax_rate_id=tax_rate.id, tax_rate_code=tax_rate.code,
                tax_rate_name=tax_rate.name, tax_rate_percent=tax_rate.rate,
                taxable_amount=taxable_amount, tax_amount=tax_amount,
            )
            db.add(sale); db.flush()
            db.add(AccountReceivable(tenant_id=actor.tenant_id, sale_id=sale.id, customer_id=counterpart.id, original_amount=total, balance_amount=total, status="open", due_date=first["due_date"] or date.today()))
            for line in lines:
                product = products[line["product_sku"]]
                db.add(SaleItem(tenant_id=actor.tenant_id, sale_id=sale.id, product_id=product.id, quantity=line["quantity"], unit_price=line["unit_value"], line_total=line["quantity"] * line["unit_value"]))
                change_stock(db, tenant_id=actor.tenant_id, product_id=product.id, warehouse_id=warehouse.id, quantity=-line["quantity"], movement_type="issue", reference=number)
            post_entry(db, tenant_id=actor.tenant_id, actor_id=actor.user_id, number=f"AS-{number}", description=f"Venta importada {number}", source_type="sale", source_id=sale.id, lines=[("1210", total, Decimal("0"), cost_center.id), ("7010", Decimal("0"), taxable_amount, cost_center.id), ("4010", Decimal("0"), tax_amount, cost_center.id)])
            record(db, tenant_id=actor.tenant_id, actor_user_id=actor.user_id, action="import", module="sales", entity="sale", entity_id=sale.id)
        else:
            purchase = Purchase(
                tenant_id=actor.tenant_id, number=number, supplier_id=counterpart.id,
                warehouse_id=warehouse.id,
                legal_entity_id=operational_context.legal_entity.id if operational_context.legal_entity else None,
                branch_id=operational_context.branch.id if operational_context.branch else None,
                status="received", total=total, note=first["note"],
                cost_center_id=cost_center.id, cost_center_code=cost_center.code,
                cost_center_name=cost_center.name,
                tax_rate_id=tax_rate.id, tax_rate_code=tax_rate.code,
                tax_rate_name=tax_rate.name, tax_rate_percent=tax_rate.rate,
                taxable_amount=taxable_amount, tax_amount=tax_amount,
            )
            db.add(purchase); db.flush()
            db.add(AccountPayable(tenant_id=actor.tenant_id, purchase_id=purchase.id, supplier_id=counterpart.id, original_amount=total, balance_amount=total, status="open", due_date=first["due_date"] or date.today()))
            for line in lines:
                product = products[line["product_sku"]]
                db.add(PurchaseItem(tenant_id=actor.tenant_id, purchase_id=purchase.id, product_id=product.id, quantity=line["quantity"], unit_cost=line["unit_value"], line_total=line["quantity"] * line["unit_value"]))
                change_stock(db, tenant_id=actor.tenant_id, product_id=product.id, warehouse_id=warehouse.id, quantity=line["quantity"], movement_type="receipt", reference=number)
            post_entry(db, tenant_id=actor.tenant_id, actor_id=actor.user_id, number=f"AS-{number}", description=f"Compra importada {number}", source_type="purchase", source_id=purchase.id, lines=[("2010", taxable_amount, Decimal("0"), cost_center.id), ("4011", tax_amount, Decimal("0"), cost_center.id), ("4210", Decimal("0"), total, cost_center.id)])
            record(db, tenant_id=actor.tenant_id, actor_user_id=actor.user_id, action="import", module="purchases", entity="purchase", entity_id=purchase.id)
    return len(grouped)


async def _preview(module: str, upload: UploadFile, actor: Actor, db: Session) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[str]]:
    if module not in TEMPLATES:
        raise HTTPException(status_code=404, detail="No hay plantilla para este módulo")
    rows = await _rows_from_upload(upload)
    valid, errors = _validate_rows(module, rows)
    conflicts = _check_database_conflicts(module, valid, actor, db) if not errors else []
    references = _transaction_references(module, valid, actor, db) if module in {"inventory", "sales", "purchases"} and not errors else []
    return valid, errors, conflicts + references


@router.get("/templates")
def list_templates(actor: Actor = Depends(require_owner_or_admin)):
    return [{"module": module, **template} for module, template in TEMPLATES.items()]


@router.post("/{module}/preview")
async def preview_import(module: str, file: UploadFile = File(...), actor: Actor = Depends(require_owner_or_admin), db: Session = Depends(get_db)):
    valid, errors, conflicts = await _preview(module, file, actor, db)
    return {"module": module, "valid_rows": len(valid), "errors": errors, "conflicts": conflicts, "can_commit": not errors and not conflicts, "sample": valid[:5]}


@router.post("/{module}/commit", status_code=status.HTTP_201_CREATED)
async def commit_import(
    module: str,
    request: Request,
    file: UploadFile = File(...),
    actor: Actor = Depends(require_owner_or_admin),
    _: Actor = Depends(require_second_factor),
    db: Session = Depends(get_db),
):
    # UploadFile is a stream.  Its digest is part of the request identity, so a
    # reused key cannot accidentally apply a different spreadsheet.  Rewind it
    # before validation because _preview deliberately reparses the source file.
    content = await file.read(MAX_FILE_SIZE + 1)
    await file.seek(0)
    # Invalid oversized files never reach a business mutation or the database.
    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(status_code=422, detail="La plantilla debe contener datos y pesar como máximo 5 MB")
    idempotency = claim_idempotency(
        db,
        tenant_id=actor.tenant_id,
        actor_user_id=actor.user_id,
        operation=f"imports.{module}.commit",
        key=request.headers.get("Idempotency-Key"),
        payload={"module": module, "filename": file.filename or "", "content_sha256": sha256(content).hexdigest()},
    )
    if idempotency and idempotency.is_replay:
        return replay_response(idempotency)
    valid, errors, conflicts = await _preview(module, file, actor, db)
    if errors or conflicts:
        raise HTTPException(status_code=422, detail={"message": "La plantilla contiene errores y no se migró ningún dato", "errors": errors, "conflicts": conflicts})
    try:
        count = _persist_master(module, valid, actor, db) if module in {"products", "customers", "suppliers", "warehouses"} else _persist_transactions(module, valid, actor, db)
        record(db, tenant_id=actor.tenant_id, actor_user_id=actor.user_id, action="bulk_import", module="imports", entity=module, detail=f"{count} registros importados desde plantilla")
        complete_idempotency(
            db,
            claim=idempotency,
            response_body={"module": module, "imported": count},
            response_status=status.HTTP_201_CREATED,
        )
        db.commit()
    except HTTPException:
        db.rollback(); raise
    except IntegrityError as exc:
        db.rollback(); raise HTTPException(status_code=409, detail="La importación entró en conflicto con datos ya existentes. No se guardó ningún registro.") from exc
    except Exception:
        db.rollback(); raise
    return {"module": module, "imported": count}
